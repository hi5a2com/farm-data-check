"""生成一份模拟的牧场繁育记录（2026 年 7–9 月），并故意埋进已知的问题。

数据全部是随机生成的，不来自任何真实牧场。埋进去的问题记在「答案」里，测试用它检查
check.py 有没有漏报或误报。

用法：
    uv run make_demo_data.py            # 写到 demo/繁育记录-演示.xlsx
    uv run make_demo_data.py -o 某处.xlsx
"""

from __future__ import annotations

import argparse
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

START, END = date(2026, 7, 1), date(2026, 9, 30)
OPERATORS = ["张敏", "李强", "王芳", "陈刚"]
TREATMENTS = ["乳房炎，头孢噻呋 3 天", "子宫炎，宫内投药", "蹄病，修蹄加药浴", "酮病，静注葡萄糖"]


def build(seed: int = 7, cows: int = 120):
    """返回（牛群名册, 繁育记录, 答案）。答案是 [(序号, 规则)]。"""
    rng = random.Random(seed)
    roster: list[dict] = []
    rows: list[dict] = []

    def add_cow() -> str:
        n = len(roster) + 1
        born = date(rng.choice([2021, 2022, 2023]), rng.randint(1, 12), rng.randint(1, 28))
        tag = f"{born.year % 100:02d}-{n:04d}"
        roster.append({"耳号": tag, "出生日期": born, "胎次": rng.randint(1, 4)})
        return tag

    def add(day: date, tag: str, event: str, detail: str) -> dict:
        row = {"日期": day, "耳号": tag, "事件": event, "详情": detail, "操作人": rng.choice(OPERATORS)}
        rows.append(row)
        return row

    def breed(day: date, tag: str) -> None:
        add(day, tag, "发情", rng.choice(["站立发情", "爬跨，黏液透明"]))
        add(day, tag, "配种", f"冻精 HO-{rng.randint(1000, 9999)}")

    def cycle(first_heat: date, tag: str) -> None:
        """从第一次发情开始配种，没怀上就约 21 天后再配，怀上了 30–35 天做孕检。"""
        h = first_heat
        while h <= END:
            breed(h, tag)
            if rng.random() < 0.6:
                check_day = h + timedelta(days=rng.randint(30, 35))
                if check_day <= END:
                    add(check_day, tag, "孕检", "B超 阳性")
                return
            h += timedelta(days=rng.randint(19, 23))

    for _ in range(cows):
        tag = add_cow()
        kind = rng.random()
        if kind < 0.4:  # 本季度产犊
            calved = START + timedelta(days=rng.randint(0, 45))
            add(calved, tag, "产犊", rng.choice(["顺产，母犊", "顺产，公犊", "助产，母犊"]))
            cycle(calved + timedelta(days=rng.randint(45, 60)), tag)
        elif kind < 0.75:  # 空怀牛
            cycle(START + timedelta(days=rng.randint(0, 20)), tag)
        if rng.random() < 0.3:
            add(START + timedelta(days=rng.randint(0, 91)), tag, "用药", rng.choice(TREATMENTS))

    answer: list[tuple[dict, str]] = []

    # 逻辑类问题：专门加几头牛
    for _ in range(3):  # W01 产后 30 天就配种
        tag = add_cow()
        calved = START + timedelta(days=rng.randint(0, 40))
        add(calved, tag, "产犊", "顺产，母犊")
        bred = calved + timedelta(days=30)
        add(bred, tag, "发情", "站立发情")
        answer.append((add(bred, tag, "配种", f"冻精 HO-{rng.randint(1000, 9999)}"), "W01"))
    for _ in range(3):  # W02 配种 20 天就孕检
        tag = add_cow()
        bred = START + timedelta(days=rng.randint(0, 60))
        breed(bred, tag)
        answer.append((add(bred + timedelta(days=20), tag, "孕检", "B超 阴性"), "W02"))
    for _ in range(2):  # W03 发情间隔 12 天
        tag = add_cow()
        h = START + timedelta(days=rng.randint(0, 60))
        breed(h, tag)
        second = add(h + timedelta(days=12), tag, "发情", "爬跨")
        add(h + timedelta(days=12), tag, "配种", f"冻精 HO-{rng.randint(1000, 9999)}")
        answer.append((second, "W03"))

    rows.sort(key=lambda r: (r["日期"], r["耳号"]))

    # 录入类问题：只改「用药」行，不影响发情、配种的逻辑检查
    planted = {id(r) for r, _ in answer}
    meds = [r for r in rows if r["事件"] == "用药"]
    rng.shuffle(meds)
    for r, bad in zip(meds[:5], ["{a}{b}", "{a}-{b_short}", "{a}-O{b3}", "{a}-{b}5", "{a}- {b}"]):
        a, b = r["耳号"].split("-")
        r["耳号"] = bad.format(a=a, b=b, b_short=b.lstrip("0"), b3=b[1:])
        answer.append((r, "E01"))
    for r, bad in zip(meds[5:9], [None, "2026/13/02", "9月5日", date(2026, 10, 20)]):
        r["日期"] = bad
        answer.append((r, "E02"))
    for r, bad in zip(meds[9:12], ["用要", "用 药", "治疗"]):
        r["事件"] = bad
        answer.append((r, "E05"))

    touched = planted | {id(r) for r in meds[:12]}
    clean = [r for r in rows if id(r) not in touched]
    rng.shuffle(clean)
    for r in clean[:4]:  # E04 操作人没填
        r["操作人"] = None
        answer.append((r, "E04"))
    for original in clean[4:8]:  # E03 同一条记录录了两遍
        dup = dict(original)
        rows.insert(rows.index(original) + 1, dup)
        answer.append((dup, "E03"))

    for i, r in enumerate(rows, start=1):
        r["序号"] = f"R{i:04d}"
    return roster, rows, [(r["序号"], rule) for r, rule in answer]


def write_xlsx(path: Path, roster: list[dict], rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "繁育记录"
    cols = ["序号", "日期", "耳号", "事件", "详情", "操作人"]
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for cell in ws["B"][1:]:
        if isinstance(cell.value, date):
            cell.number_format = "yyyy-mm-dd"
    herd = wb.create_sheet("牛群名册")
    herd.append(["耳号", "出生日期", "胎次"])
    for c in roster:
        herd.append([c["耳号"], c["出生日期"], c["胎次"]])
    for cell in herd["B"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for sheet in (ws, herd):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="生成模拟繁育记录")
    parser.add_argument("-o", "--output", type=Path, default=Path("demo/繁育记录-演示.xlsx"))
    parser.add_argument("--seed", type=int, default=7)
    args = parser.parse_args()
    roster, rows, answer = build(args.seed)
    write_xlsx(args.output, roster, rows)
    print(f"已生成 {len(roster)} 头牛、{len(rows)} 条记录，埋入 {len(answer)} 处问题：{args.output}")


if __name__ == "__main__":
    main()
