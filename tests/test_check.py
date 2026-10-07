from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import check
import make_demo_data

TODAY = date(2026, 10, 6)


def run(tmp_path: Path, seed: int):
    roster, rows, answer = make_demo_data.build(seed)
    src = tmp_path / "繁育记录.xlsx"
    make_demo_data.write_xlsx(src, roster, rows)
    records, herd = check.load_records(src)
    issues = check.check(records, herd, TODAY)
    return src, records, issues, answer


@pytest.mark.parametrize("seed", [7, 1, 2, 3, 42])
def test_finds_every_planted_problem_and_nothing_else(tmp_path, seed):
    _, _, issues, answer = run(tmp_path, seed)
    assert {(i.no, i.rule) for i in issues} == set(answer)


def test_every_rule_is_exercised(tmp_path):
    _, _, issues, _ = run(tmp_path, 7)
    assert {i.rule for i in issues} == set(check.RULES)


def test_clean_data_has_no_issues(tmp_path):
    roster, rows, answer = make_demo_data.build(7)
    bad = {no for no, _ in answer}
    dup_of = {no for no, rule in answer if rule == "E03"}
    clean_rows = [r for r in rows if r["序号"] not in bad and r["序号"] not in dup_of]
    # 去掉专门制造逻辑问题的那几头牛，剩下的应该完全干净
    special = {r["耳号"] for r in rows if r["序号"] in bad and r["事件"] in ("配种", "孕检", "发情")}
    clean_rows = [r for r in clean_rows if r["耳号"] not in special]
    src = tmp_path / "干净.xlsx"
    make_demo_data.write_xlsx(src, roster, clean_rows)
    records, herd = check.load_records(src)
    assert check.check(records, herd, TODAY) == []


def test_report_marks_problem_rows(tmp_path):
    src, records, issues, _ = run(tmp_path, 7)
    out = tmp_path / "结果.xlsx"
    check.write_report(src, out, records, issues)
    wb = load_workbook(out)
    assert wb.sheetnames == ["汇总", "问题清单", "原表标注"]
    assert wb["问题清单"].max_row == len(issues) + 1
    marked = wb["原表标注"]
    problem_rows = {r.row for r in records if r.no in {i.no for i in issues}}
    filled = {c.row for c in marked["A"] if c.fill.fgColor.rgb not in (None, "00000000")}
    assert filled == problem_rows
